"""
ファイル出力ユーティリティ

CSV、PDF、DOCX、Markdownなどの形式でファイルを出力する機能を提供
"""
import csv
import io
import json
import logging
from typing import Optional, Dict, Any, List
from pathlib import Path
import tempfile
import os

logger = logging.getLogger(__name__)


def export_to_csv(content: str, filename: Optional[str] = None) -> tuple:
    """
    テキストコンテンツをCSV形式で出力

    Args:
        content: 出力するテキストコンテンツ
        filename: ファイル名（省略時は自動生成）

    Returns:
        (ファイルのバイトデータ, ファイル名)
    """
    try:
        # テキストをCSV形式に変換
        # まず、テキストを行に分割
        lines = content.strip().split('\n')

        # CSVライターで出力
        output = io.StringIO()
        writer = csv.writer(output)

        # 各行を処理
        for line in lines:
            # タブ区切りの場合はそのまま、そうでなければカンマ区切りとして処理
            if '\t' in line:
                writer.writerow(line.split('\t'))
            elif ',' in line and not line.startswith('#'):
                # カンマ区切りの場合（ただし、コメント行は除外）
                writer.writerow([cell.strip() for cell in line.split(',')])
            else:
                # 通常のテキスト行
                writer.writerow([line])

        csv_content = output.getvalue()
        output.close()

        # バイトデータに変換
        csv_bytes = csv_content.encode('utf-8-sig')  # BOM付きUTF-8（Excel対応）

        if not filename:
            filename = "output.csv"
        elif not filename.endswith('.csv'):
            filename = f"{filename}.csv"

        return csv_bytes, filename

    except Exception as e:
        logger.error(f"CSV出力エラー: {str(e)}")
        # フォールバック: テキストをそのままCSVとして出力
        csv_bytes = content.encode('utf-8-sig')
        filename = filename or "output.csv"
        return csv_bytes, filename


def export_to_markdown(content: str, filename: Optional[str] = None) -> tuple:
    """
    テキストコンテンツをMarkdown形式で出力

    Args:
        content: 出力するテキストコンテンツ
        filename: ファイル名（省略時は自動生成）

    Returns:
        (ファイルのバイトデータ, ファイル名)
    """
    md_bytes = content.encode('utf-8')

    if not filename:
        filename = "output.md"
    elif not filename.endswith('.md'):
        filename = f"{filename}.md"

    return md_bytes, filename


def export_to_txt(content: str, filename: Optional[str] = None) -> tuple:
    """
    テキストコンテンツをTXT形式で出力

    Args:
        content: 出力するテキストコンテンツ
        filename: ファイル名（省略時は自動生成）

    Returns:
        (ファイルのバイトデータ, ファイル名)
    """
    txt_bytes = content.encode('utf-8')

    if not filename:
        filename = "output.txt"
    elif not filename.endswith('.txt'):
        filename = f"{filename}.txt"

    return txt_bytes, filename


def export_to_pdf(content: str, filename: Optional[str] = None) -> tuple:
    """
    テキストコンテンツをPDF形式で出力

    Args:
        content: 出力するテキストコンテンツ
        filename: ファイル名（省略時は自動生成）

    Returns:
        (ファイルのバイトデータ, ファイル名)

    Note:
        PDF生成にはreportlabが必要です。インストールされていない場合は
        テキスト形式で返します。
    """
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        # PDF生成
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=20*mm,
            leftMargin=20*mm,
            topMargin=20*mm,
            bottomMargin=20*mm
        )

        # 日本語フォントの登録を試行（利用可能な場合）
        japanese_font_registered = False
        try:
            # システムに存在する可能性のある日本語フォントを試行
            font_paths = [
                '/System/Library/Fonts/ヒラギノ角ゴシック W3.ttc',  # macOS
                '/System/Library/Fonts/ヒラギノ角ゴシック W6.ttc',  # macOS
                '/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc',  # Linux
                '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',  # Linux
            ]
            for font_path in font_paths:
                if os.path.exists(font_path):
                    try:
                        pdfmetrics.registerFont(TTFont('JapaneseFont', font_path))
                        japanese_font_registered = True
                        logger.info(f"日本語フォントを登録しました: {font_path}")
                        break
                    except Exception as font_error:
                        logger.warning(f"フォントファイルの読み込みに失敗しました ({font_path}): {str(font_error)}")
                        continue
        except Exception as e:
            logger.warning(f"日本語フォントの登録に失敗しました: {str(e)}")

        # スタイル設定
        styles = getSampleStyleSheet()
        normal_style = ParagraphStyle(
            'CustomNormal',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            fontName='JapaneseFont' if japanese_font_registered else 'Helvetica',
        )

        # コンテンツを段落に分割
        story = []
        lines = content.split('\n')

        for line in lines:
            if line.strip():
                # HTMLエスケープ処理（日本語対応）
                line_escaped = line.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                story.append(Paragraph(line_escaped, normal_style))
            else:
                story.append(Spacer(1, 6))

        # PDF生成
        doc.build(story)
        pdf_bytes = buffer.getvalue()
        buffer.close()

        if not filename:
            filename = "output.pdf"
        elif not filename.endswith('.pdf'):
            filename = f"{filename}.pdf"

        return pdf_bytes, filename

    except ImportError:
        logger.warning("reportlabがインストールされていません。PDF出力は利用できません。")
        # フォールバック: テキスト形式で返す
        return export_to_txt(content, filename)
    except Exception as e:
        logger.error(f"PDF出力エラー: {str(e)}")
        # フォールバック: テキスト形式で返す
        return export_to_txt(content, filename)


def export_to_docx(content: str, filename: Optional[str] = None) -> tuple:
    """
    テキストコンテンツをDOCX形式で出力

    Args:
        content: 出力するテキストコンテンツ
        filename: ファイル名（省略時は自動生成）

    Returns:
        (ファイルのバイトデータ, ファイル名)

    Note:
        DOCX生成にはpython-docxが必要です。インストールされていない場合は
        テキスト形式で返します。
    """
    try:
        from docx import Document
        from docx.shared import Pt

        # ドキュメント作成
        doc = Document()

        # スタイル設定（日本語フォント対応）
        style = doc.styles['Normal']
        font = style.font
        # 日本語フォントを試行（存在しない場合はシステムのデフォルトフォントを使用）
        japanese_fonts = ['游ゴシック', 'Yu Gothic', 'MS Gothic', 'Hiragino Sans', 'Noto Sans CJK JP']
        font.name = japanese_fonts[0]  # 最初のフォントを試行
        font.size = Pt(10)

        # コンテンツを行ごとに追加
        lines = content.split('\n')
        for line in lines:
            if line.strip():
                # 見出し判定（#で始まる行）
                if line.startswith('# '):
                    doc.add_heading(line[2:].strip(), level=1)
                elif line.startswith('## '):
                    doc.add_heading(line[3:].strip(), level=2)
                elif line.startswith('### '):
                    doc.add_heading(line[4:].strip(), level=3)
                else:
                    doc.add_paragraph(line)
            else:
                doc.add_paragraph('')

        # メモリに保存
        buffer = io.BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()
        buffer.close()

        if not filename:
            filename = "output.docx"
        elif not filename.endswith('.docx'):
            filename = f"{filename}.docx"

        return docx_bytes, filename

    except ImportError:
        logger.warning("python-docxがインストールされていません。DOCX出力は利用できません。")
        # フォールバック: テキスト形式で返す
        return export_to_txt(content, filename)
    except Exception as e:
        logger.error(f"DOCX出力エラー: {str(e)}")
        # フォールバック: テキスト形式で返す
        return export_to_txt(content, filename)


def export_content(
    content: str,
    output_format: str = "txt",
    filename: Optional[str] = None
) -> tuple:
    """
    コンテンツを指定された形式で出力

    Args:
        content: 出力するテキストコンテンツ
        output_format: 出力形式（csv, pdf, docx, md, txt）
        filename: ファイル名（省略時は自動生成）

    Returns:
        (ファイルのバイトデータ, ファイル名)

    Raises:
        ValueError: サポートされていない形式が指定された場合
    """
    format_lower = output_format.lower()

    if format_lower == "csv":
        return export_to_csv(content, filename)
    elif format_lower == "pdf":
        return export_to_pdf(content, filename)
    elif format_lower == "docx":
        return export_to_docx(content, filename)
    elif format_lower in ["md", "markdown"]:
        return export_to_markdown(content, filename)
    elif format_lower == "txt":
        return export_to_txt(content, filename)
    else:
        raise ValueError(f"サポートされていない出力形式: {output_format}")


def parse_attached_files(attachments: Optional[List[Dict[str, Any]]] = None) -> Dict[str, str]:
    """
    添付ファイルを解析して、ファイル名と内容の辞書を返す

    Args:
        attachments: 添付ファイルのリスト（各要素は{"filename": str, "content": str}形式）

    Returns:
        ファイル名をキー、内容を値とする辞書
    """
    if not attachments:
        return {}

    result = {}
    for attachment in attachments:
        filename = attachment.get("filename", "")
        content = attachment.get("content", "")
        if filename and content:
            result[filename] = content

    return result

